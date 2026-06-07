import os
from flask import Flask, render_template, request, redirect, url_for
from flask_sqlalchemy import SQLAlchemy

app = Flask(__name__)

basedir = os.path.abspath(os.path.dirname(__file__))
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(basedir, 'books.db')
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {
    "connect_args": {"check_same_thread": False, "uri": True}
}
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

# 定義資料庫欄位
class Book(db.Model):
    id = db.Column(db.Integer, primary_key=True)       
    title = db.Column(db.String(100), nullable=False)   
    author = db.Column(db.String(50), nullable=False)   
    isbn = db.Column(db.String(20), nullable=True)      
    publisher = db.Column(db.String(120))               
    quantity = db.Column(db.Integer, default=1)         
    translator = db.Column(db.String(50))               
    status = db.Column(db.String(20), nullable=False, default='未完成閱讀') 
    category = db.Column(db.String(200), nullable=True) # 書籍類別欄位 (儲存如 "心理學,外語書")

with app.app_context():
    db.create_all()

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        book_title = request.form.get('title')
        book_author = request.form.get('author')
        book_isbn = request.form.get('isbn') or None
        book_publisher = request.form.get('publisher') or None    
        book_quantity = request.form.get('quantity')      
        book_translator = request.form.get('translator') or None  
        
        # 獲取前端複選框的陣列資料
        selected_categories = request.form.getlist('category')
        # 用逗號將多個類別串接起來，若沒勾選則存入 '-'
        book_category = ",".join(selected_categories) if selected_categories else '-'

        new_book = Book(
            title=book_title, 
            author=book_author, 
            translator=book_translator,  
            isbn=book_isbn,
            publisher=book_publisher,
            quantity=int(book_quantity) if book_quantity else 1,
            status='未完成閱讀',
            category=book_category 
        )
        db.session.add(new_book)
        db.session.commit()
        return redirect(url_for('index'))
        
    search_query = request.args.get('search', '')
    
    if search_query:
        from sqlalchemy import or_
        all_books = Book.query.filter(
            or_(
                Book.title.like(f"%{search_query}%"),
                Book.author.like(f"%{search_query}%"),
                Book.publisher.like(f"%{search_query}%"),
                Book.isbn.like(f"%{search_query}%"),
                Book.translator.like(f"%{search_query}%"),
                Book.status.like(f"%{search_query}%"),
                Book.category.like(f"%{search_query}%") # 支援搜尋書籍類別（如搜尋：外語書）
            )
        ).all()
    else:
        all_books = Book.query.all()
        
    return render_template('index.html', books=all_books, search_query=search_query)

@app.route('/toggle_status/<int:book_id>')
def toggle_status(book_id):
    book = Book.query.get_or_404(book_id)
    if book.status == '未完成閱讀':
        book.status = '已完成閱讀'
    else:
        book.status = '未完成閱讀'
    db.session.commit()
    return redirect(url_for('index'))

@app.route('/delete/<int:book_id>')
def delete_book(book_id):
    book_to_delete = Book.query.get_or_404(book_id)
    db.session.delete(book_to_delete)
    db.session.commit()
    return redirect(url_for('index'))

@app.route('/export_excel')
def export_excel():
    import pandas as pd
    import io
    from flask import Response
    from datetime import datetime
    import openpyxl
    from openpyxl.styles import Font, Alignment

    all_books = Book.query.all()
    
    books_data = []
    for b in all_books:
        books_data.append({
            "書籍 ID": b.id,
            "書名": b.title,
            "作者": b.author,
            "譯者": b.translator or '-',
            "出版社": b.publisher or '-',
            "ISBN": b.isbn or '-',
            "書籍類別": b.category or '-', 
            "數量": b.quantity,
            "未完成閱讀": "V" if b.status == "未完成閱讀" else "", 
            "已完成閱讀": "V" if b.status == "已完成閱讀" else ""  
        })
    
    df = pd.DataFrame(books_data)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "館藏清單"
    ws.views.sheetView[0].showGridLines = True

    ws['A1'] = "Gary的圖書管理資料"
    ws['A1'].font = Font(size=16, bold=True)

    # 日期對齊 J1 (第 10 欄)
    today_str = f"匯出日期: {datetime.now().strftime('%Y-%m-%d')}"
    ws['J1'] = today_str
    ws['J1'].font = Font(size=10, italic=True)
    ws['J1'].alignment = Alignment(horizontal='right')

    headers = list(df.columns)
    ws.append([]) 
    ws.append(headers) 

    for row in df.values.tolist():
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment;filename=book_list.xlsx"}
    )

if __name__ == '__main__':
    app.run(debug=True, port=1996)